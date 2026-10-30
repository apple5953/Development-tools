# -*- coding: utf-8 -*-
"""
釋疑簡報(預計提出前的問題紀錄) 產生器
用法：
  python gen_rfi_deck.py 輸入.xlsx [輸出.pptx]      ← 同事最常用：填 Excel 清單
  python gen_rfi_deck.py 輸入.json [輸出.pptx]      ← AI 整理後直接給 JSON

每一題產生：
  ‧ 問題紀錄頁：問題紀錄-預計提出Q001 / 圖說編號名稱 / 位置 / 問題 / 建議 / 回覆歷程 + 右側圖說(紅字圖名)
  ‧ (選用) 調整建議頁：調整說明 + 調整後示意圖
整份另含：封面、統計圓餅圖(類別、已解決/未解決)、問題總表、分段頁(RFI-003…)、END 頁。
欄位定義見 工作流程/03_釋疑簡報.md
"""
import json
import math
import os
import sys
from collections import Counter, OrderedDict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt
from pptx_tools import add_picture_fit, add_red_box, emu

FONT = "微軟正黑體"
C_TEXT = RGBColor(0x26, 0x26, 0x26)
C_BLUE = RGBColor(0x1F, 0x5F, 0xBF)     # 建議
C_GREEN = RGBColor(0x2E, 0x8B, 0x3E)    # 回覆歷程
C_RED = RGBColor(0xE0, 0x30, 0x30)      # 圖名 / 強調
C_GREY = RGBColor(0x8A, 0x8A, 0x8A)
C_LINE = RGBColor(0xD9, 0xD9, 0xD9)
CATEGORY = OrderedDict([("H", "H圖說釋疑"), ("B", "B結構設計"), ("C", "C建築設計"), ("D", "D介面衝突")])
CAT_COLOR = {"H": RGBColor(0xA2, 0x2E, 0x8C), "B": RGBColor(0x5B, 0xD6, 0xE0),
             "C": RGBColor(0x7F, 0xC8, 0xF5), "D": RGBColor(0x6F, 0xA8, 0xDC)}
STATUS_COLOR = {"預計提出": RGBColor(0xF2, 0x9B, 0x1E), "已提出": RGBColor(0x2E, 0x8B, 0x3E),
                "暫不提出": RGBColor(0x8A, 0x8A, 0x8A), "不提出": RGBColor(0x8A, 0x8A, 0x8A)}

COLS = ["編號", "群組", "類別", "狀態", "RFI編號", "是否解決", "圖說編號名稱", "位置", "問題", "建議", "回覆歷程",
        "圖片1", "圖說1", "圖片2", "圖說2", "圖片3", "圖說3", "圖片4", "圖說4",
        "調整說明", "調整圖片1", "調整圖說1", "調整圖片2", "調整圖說2", "備註",
        "其他圖片", "其他圖說", "追蹤編號", "提出日期", "回覆期限", "事務所回覆", "定案結論", "定案日期", "追蹤狀態", "來源", "Revit視圖", "元件ID"]
TRACK_COLOR = {"預計提出": RGBColor(0xF2, 0x9B, 0x1E), "待回覆": RGBColor(0xD9, 0x3A, 0x3A),
               "已回覆待確認": RGBColor(0x2F, 0x75, 0xB5), "已定案": RGBColor(0x2E, 0x8B, 0x3E),
               "不提出": RGBColor(0x8A, 0x8A, 0x8A)}


def _cell_str(v):
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y/%m/%d")
    return str(v).strip()

# ------------------------------------------------------------------ input
def load_input(path):
    base = os.path.dirname(os.path.abspath(path))
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        meta = {}
        if "設定" in wb.sheetnames:
            for k, v in wb["設定"].iter_rows(min_row=1, values_only=True):
                if k:
                    meta[str(k).strip()] = "" if v is None else str(v)
        ws = wb["問題清單"] if "問題清單" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() if h else "" for h in rows[0]]
        items = []
        for r in rows[1:]:
            rec = {head[i]: _cell_str(v) for i, v in enumerate(r) if i < len(head)}
            if not rec.get("編號") and not rec.get("問題"):
                continue
            items.append(_from_flat(rec))
        cfg = {"project": meta.get("專案名稱", ""), "title": meta.get("簡報標題", "BIM釋疑-預計提出"),
               "subtitle": meta.get("副標題", ""), "date": meta.get("日期", ""),
               "company": meta.get("製作單位", ""), "output": meta.get("輸出檔名", ""),
               "show_stats": meta.get("顯示統計", "是") != "否", "show_index": meta.get("顯示總表", "是") != "否",
               "items": items}
    else:
        cfg = json.load(open(path, encoding="utf-8"))
        cfg["items"] = [_from_flat(x) if "images" not in x else x for x in cfg.get("items", [])]
    for it in cfg["items"]:
        for im in it.get("images", []) + it.get("adjust_images", []):
            if im.get("path") and not os.path.isabs(im["path"]):
                im["path"] = os.path.normpath(os.path.join(base, im["path"]))
    return cfg


def _from_flat(r):
    imgs = [{"path": r.get(f"圖片{i}"), "caption": r.get(f"圖說{i}", "")} for i in range(1, 5) if r.get(f"圖片{i}")]
    more = [x.strip() for x in (r.get("其他圖片") or "").split(";") if x.strip()]
    mcap = (r.get("其他圖說") or "").split(";")
    imgs += [{"path": pth, "caption": (mcap[i].strip() if i < len(mcap) else "")} for i, pth in enumerate(more)]
    adj = [{"path": r.get(f"調整圖片{i}"), "caption": r.get(f"調整圖說{i}", "")} for i in range(1, 3) if r.get(f"調整圖片{i}")]
    return {"no": r.get("編號", ""), "group": r.get("群組", ""), "category": (r.get("類別", "") or "")[:1].upper(),
            "status": r.get("狀態", "") or "預計提出", "rfi_no": r.get("RFI編號", ""),
            "solved": r.get("是否解決", ""), "drawing": r.get("圖說編號名稱", ""), "location": r.get("位置", ""),
            "question": r.get("問題", ""), "suggest": r.get("建議", ""), "replies": r.get("回覆歷程", ""),
            "images": imgs, "adjust": r.get("調整說明", ""), "adjust_images": adj, "note": r.get("備註", ""),
            "raised": r.get("提出日期", ""), "due": r.get("回覆期限", ""), "office_reply": r.get("事務所回覆", ""),
            "decision": r.get("定案結論", ""), "decided": r.get("定案日期", ""), "track": r.get("追蹤狀態", ""),
            "source": r.get("來源", ""), "tid": r.get("追蹤編號", "")}

# ------------------------------------------------------------------ drawing helpers
def tb(slide, x, y, w, h, paras, size=13, anchor=MSO_ANCHOR.TOP, align=None, autofit=False):
    """paras: [[(text, {color,bold,size})...], ...]"""
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        p.space_after = Pt(2)
        if isinstance(runs, str):
            runs = [(runs, {})]
        for text, st in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", False)
            f.color.rgb = st.get("color", C_TEXT)
            rPr = r._r.get_or_add_rPr()
            ea = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": FONT})
            rPr.append(ea)
    return box


def chip(slide, x, y, text, fill, w=None):
    w = w or max(2.2, 0.45 * len(text) + 0.6)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(0.75))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_top = tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return w


def line(slide, x1, y1, x2, y2, color=C_LINE, pt=1):
    c = slide.shapes.add_connector(1, emu(x1), emu(y1), emu(x2), emu(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)


def footer(slide, cfg, n):
    tb(slide, 0.9, 18.25, 20, 0.6, [[(f'{cfg.get("project", "")}　{cfg.get("title", "")}', {"size": 9, "color": C_GREY})]])
    tb(slide, 30.5, 18.25, 2.5, 0.6, [[(str(n), {"size": 9, "color": C_GREY})]], align=PP_ALIGN.RIGHT)


def est_h(text, cpl, lh=0.56):
    if not text:
        return 0
    return sum(max(1, math.ceil(len(l) / cpl)) for l in str(text).split("\n")) * lh


MAIN_MAX, MORE_MAX = 2, 4     # 問題頁最多 2 張圖；其餘放「續頁」，每頁最多 4 張


def _img_size(path):
    try:
        from PIL import Image
        with Image.open(path) as im:
            return im.size
    except Exception:
        return (16, 9)


def best_layout(sizes, w, h, cap, gap):
    """依各圖長寬比，在 w×h 範圍找讓圖片總面積最大的排法（列數×欄數）。
    回傳 [(cx, cy, cw, ch)]：每張圖的格子（相對範圍左上角，ch 已扣掉圖說高度）"""
    n = len(sizes)
    best, best_score = None, -1
    for rows in range(1, n + 1):
        cols = math.ceil(n / rows)
        if (rows - 1) * cols >= n:          # 有整列空著 → 不考慮
            continue
        cw = (w - gap * (cols - 1)) / cols
        ch = (h - gap * (rows - 1)) / rows - cap
        if cw <= 0.5 or ch <= 0.5:
            continue
        score = 0
        for iw, ih in sizes:
            k = min(cw / iw, ch / ih)
            score += k * iw * k * ih
        if score > best_score:
            best_score, best = score, (rows, cols, cw, ch)
    rows, cols, cw, ch = best
    cells = []
    for i in range(n):
        r, c = divmod(i, cols)
        in_row = min(cols, n - r * cols)
        offset = (cols - in_row) * (cw + gap) / 2      # 最後一列置中
        cells.append((offset + c * (cw + gap), r * (ch + cap + gap), cw, ch))
    return cells


def image_grid(slide, imgs, x, y, w, h):
    n = len(imgs)
    if not n:
        return
    gap, cap = 0.35, 0.75
    sizes = [_img_size(i["path"]) if i.get("path") and os.path.exists(i["path"]) else (16, 9) for i in imgs]
    for im, (dx, dy, cw, ch) in zip(imgs, best_layout(sizes, w, h, cap, gap)):
        cx, cy = x + dx, y + dy
        if not im.get("path") or not os.path.exists(im["path"]):
            ph = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(cx), emu(cy), emu(cw), emu(ch))
            ph.fill.solid(); ph.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
            ph.line.color.rgb = C_LINE
            ph.text_frame.text = "（圖片待補：%s）" % os.path.basename(im.get("path") or "")
            for p_ in ph.text_frame.paragraphs:
                for r_ in p_.runs:
                    r_.font.size = Pt(12); r_.font.name = FONT; r_.font.color.rgb = C_GREY
            pic_l, pic_t, pic_w, pic_h = emu(cx), emu(cy), emu(cw), emu(ch)
        else:
            pic = add_picture_fit(slide, im["path"], emu(cx), emu(cy), emu(cw), emu(ch))
            pic.line.color.rgb = C_LINE
            pic.line.width = Pt(0.75)
            pic_l, pic_t, pic_w, pic_h = pic.left, pic.top, pic.width, pic.height
            for mk in im.get("marks", []):
                add_red_box(slide, int(pic_l + mk[0] * pic_w), int(pic_t + mk[1] * pic_h),
                            int(mk[2] * pic_w), int(mk[3] * pic_h))
        if im.get("caption"):
            tb(slide, cx, (pic_t + pic_h) / 360000 + 0.05, cw, cap,
               [[(im["caption"], {"size": 12, "bold": True, "color": C_RED})]], align=PP_ALIGN.CENTER)


def split_even(items, per):
    """平均分頁：7 張、每頁最多 4 → 4+3；5 張 → 3+2"""
    if not items:
        return []
    k = math.ceil(len(items) / per)
    base, extra = divmod(len(items), k)
    out, i = [], 0
    for j in range(k):
        n = base + (1 if j < extra else 0)
        out.append(items[i:i + n])
        i += n
    return out


def slide_more(prs, cfg, it, imgs, k, total, n):
    """圖片續頁：只有圖，不重複文字"""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    head = cfg.get("q_prefix") or ("問題紀錄-預計提出" if it.get("status", "預計提出") == "預計提出" else "問題紀錄 ")
    tb(s, 0.7, 0.35, 22, 1.3, [[(f'{head}{it.get("no", "")}', {"size": 24}),
                                (f"　（續 {k}/{total}）", {"size": 16, "color": C_GREY})]])
    line(s, 0.7, 1.75, 32.9, 1.75, C_LINE, 1)
    image_grid(s, imgs, 0.9, 2.1, 32.0, 15.9)
    footer(s, cfg, n)


# ------------------------------------------------------------------ slides
def slide_cover(prs, cfg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, emu(7.2), prs.slide_width, emu(0.18))
    band.fill.solid(); band.fill.fore_color.rgb = C_RED; band.line.fill.background()
    tb(s, 2, 3.2, 30, 3.2, [[(cfg.get("title", "BIM釋疑-預計提出"), {"size": 54, "bold": True})]], anchor=MSO_ANCHOR.BOTTOM)
    paras = [[(cfg.get("project", ""), {"size": 22, "bold": True})]]
    if cfg.get("subtitle"):
        for l in cfg["subtitle"].split("\n"):
            paras.append([(l, {"size": 18, "color": C_GREY})])
    tb(s, 2, 8.0, 30, 5, paras)
    tb(s, 2, 15.6, 30, 1.6, [[(cfg.get("company", ""), {"size": 16, "bold": True})],
                             [(cfg.get("date", ""), {"size": 14, "color": C_GREY})]])
    return s


def slide_big(prs, cfg, text, sub="", n=0):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, 1, 6.2, 31.9, 4.5, [[(text, {"size": 80, "bold": True})]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    if sub:
        tb(s, 1, 10.9, 31.9, 1.5, [[(sub, {"size": 24, "color": C_GREY})]], align=PP_ALIGN.CENTER)
    if n:
        footer(s, cfg, n)
    return s


def _pie(slide, x, y, w, h, title, cats, vals, colors):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(title, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.PIE, emu(x), emu(y), emu(w), emu(h), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    ch.legend.font.name = FONT
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.show_value = True
    dl.show_category_name = False
    dl.number_format = '0"題"'
    dl.number_format_is_linked = False
    dl.font.size = Pt(14)
    dl.font.bold = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, c in enumerate(colors):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c


def slide_stats(prs, cfg, items, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, 0.9, 0.5, 20, 1.3, [[("釋疑統計", {"size": 26, "bold": True})]])
    line(s, 0.9, 1.8, 32.9, 1.8, C_RED, 1.5)
    cc = Counter(i.get("category") or "未分類" for i in items)
    cats = [k for k in CATEGORY if cc.get(k)] + (["未分類"] if cc.get("未分類") else [])
    if cats:
        _pie(s, 1.0, 2.3, 15.5, 15.2, "釋疑累積圖(依類別區分)", [CATEGORY.get(k, k) for k in cats],
             [cc[k] for k in cats], [CAT_COLOR.get(k, C_GREY) for k in cats])
    solved = sum(1 for i in items if str(i.get("solved", "")).strip() in ("是", "已解決", "Y", "y", "1", "True"))
    _pie(s, 17.3, 2.3, 15.5, 15.2, f"總釋疑累積圖（共 {len(items)} 題）", ["已解決", "未解決"],
         [solved, len(items) - solved], [RGBColor(0x7E, 0xD3, 0x57), RGBColor(0x6E, 0xD3, 0xF5)])
    footer(s, cfg, n)


def slide_index(prs, cfg, items, start_n):
    per = 12
    pages = math.ceil(len(items) / per) or 1
    for pg in range(pages):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb(s, 0.9, 0.5, 25, 1.3, [[("問題總表" + (f"（{pg + 1}/{pages}）" if pages > 1 else ""), {"size": 26, "bold": True})]])
        line(s, 0.9, 1.8, 32.9, 1.8, C_RED, 1.5)
        chunk = items[pg * per:(pg + 1) * per]
        hdr = ["編號", "類別", "位置", "問題摘要", "狀態"]
        widths = [2.2, 3.0, 6.5, 16.6, 3.7]
        t = s.shapes.add_table(len(chunk) + 1, 5, emu(0.9), emu(2.3), emu(sum(widths)), emu(1.2 * (len(chunk) + 1))).table
        for i, wd in enumerate(widths):
            t.columns[i].width = emu(wd)
        for r in range(len(chunk) + 1):
            t.rows[r].height = emu(1.15)
            vals = hdr if r == 0 else [chunk[r - 1].get("no", ""), CATEGORY.get(chunk[r - 1].get("category"), ""),
                                       chunk[r - 1].get("location", ""),
                                       (chunk[r - 1].get("question", "").replace("\n", " "))[:60],
                                       chunk[r - 1].get("status", "") + (("\n" + chunk[r - 1]["rfi_no"]) if chunk[r - 1].get("rfi_no") else "")]
            for c, v in enumerate(vals):
                cell = t.cell(r, c)
                cell.text = v
                cell.margin_left = cell.margin_right = Cm(0.12)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(11 if r else 12)
                        run.font.name = FONT
                        run.font.bold = (r == 0)
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else C_TEXT
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x40, 0x40, 0x40) if r == 0 else (
                    RGBColor(0xF7, 0xF7, 0xF7) if r % 2 else RGBColor(0xFF, 0xFF, 0xFF))
        footer(s, cfg, start_n + pg)
    return pages


def slide_question(prs, cfg, it, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    head = cfg.get("q_prefix") or ("問題紀錄-預計提出" if it.get("status", "預計提出") == "預計提出" else "問題紀錄 ")
    tb(s, 0.7, 0.35, 16, 1.3, [[(f'{head}{it.get("no", "")}', {"size": 24})]])
    # 右上標籤：類別 / 狀態
    x = 32.9
    st = it.get("status", "")
    if it.get("rfi_no"):
        st = f'{st} {it["rfi_no"]}'.strip()
    if st:
        w = max(2.2, 0.45 * len(st) + 0.6); x -= w
        chip(s, x, 0.55, st, STATUS_COLOR.get(it.get("status", ""), C_GREY), w)
    if it.get("category") in CATEGORY:
        w = 3.0; x -= w + 0.25
        chip(s, x, 0.55, CATEGORY[it["category"]], CAT_COLOR[it["category"]], w)
    tr = it.get("track", "")
    if cfg.get("show_track", True) and tr and tr != it.get("status") and tr in TRACK_COLOR:
        w = max(2.2, 0.45 * len(tr) + 0.6); x -= w + 0.25
        chip(s, x, 0.55, tr, TRACK_COLOR[tr], w)
    info = []
    if it.get("tid") and cfg.get("show_tid", True):
        info.append(f'追蹤編號 {it["tid"]}')
    if it.get("raised"):
        info.append(f'提出 {it["raised"]}')
        if it.get("due") and tr in ("待回覆", ""):
            info.append(f'回覆期限 {it["due"]}')
    if info:
        tb(s, 16.9, 1.15, x - 17.1, 0.6, [[("　".join(info), {"size": 10, "color": C_GREY})]], align=PP_ALIGN.RIGHT)
    line(s, 0.7, 1.75, 32.9, 1.75, C_LINE, 1)

    # 左欄文字
    LX, LW, cpl = 0.9, 12.0, 21
    y = 2.1
    paras = [[("圖說編號名稱：", {"bold": True})]]
    for l in str(it.get("drawing", "")).split("\n"):
        paras.append([(l, {})])
    paras.append([("位置：", {"bold": True}), (it.get("location", ""), {})])
    paras.append([("問題：", {"bold": True})])
    for l in str(it.get("question", "")).split("\n"):
        paras.append([(l, {})])
    h = est_h(it.get("drawing", ""), cpl) + est_h(it.get("question", ""), cpl) + 0.56 * 3 + 0.3
    tb(s, LX, y, LW, h, paras)
    y += h + 0.25
    if it.get("suggest"):
        sp = [[("建議：", {"bold": True, "color": C_BLUE})]] + \
             [[(l, {"color": C_BLUE})] for l in str(it["suggest"]).split("\n")]
        h = est_h(it["suggest"], cpl) + 0.56 + 0.2
        tb(s, LX, y, LW, h, sp)
        y += h + 0.25
    if it.get("replies"):
        rp = [[("回覆歷程：", {"bold": True, "color": C_GREEN, "size": 12})]] + \
             [[(l, {"color": C_GREEN, "size": 12})] for l in str(it["replies"]).split("\n")]
        h = est_h(it["replies"], cpl + 2, 0.52) + 0.52 + 0.2
        tb(s, LX, y, LW, h, rp, size=12)
        y += h
    if it.get("office_reply") and it.get("office_reply") not in str(it.get("replies", "")):
        rp = [[("事務所回覆：", {"bold": True, "color": C_GREEN})]] + \
             [[(l, {"color": C_GREEN})] for l in str(it["office_reply"]).split("\n")]
        h = est_h(it["office_reply"], cpl) + 0.56 + 0.2
        tb(s, LX, y + 0.15, LW, h, rp)
        y += h + 0.25
    if it.get("decision"):
        dp = [[("定案" + (f'（{it["decided"]}）' if it.get("decided") else "") + "：", {"bold": True, "color": C_RED})]] + \
             [[(l, {"bold": True, "color": C_RED})] for l in str(it["decision"]).split("\n")]
        h = est_h(it["decision"], cpl) + 0.56 + 0.2
        tb(s, LX, y + 0.15, LW, h, dp)
        y += h + 0.25
    if it.get("note"):
        note = it["note"] if len(it["note"]) <= 90 else it["note"][:88] + "…"
        tb(s, LX, max(y + 0.2, 16.4), LW, 1.4, [[("備註：" + note, {"size": 10, "color": C_GREY})]])
    line(s, 13.3, 2.1, 13.3, 17.9, C_LINE, 0.75)
    imgs = it.get("images", [])
    image_grid(s, imgs[:MAIN_MAX], 13.7, 2.1, 19.2, 15.9)
    footer(s, cfg, n)
    rest = imgs[MAIN_MAX:]
    chunks = split_even(rest, MORE_MAX)
    for k, ch in enumerate(chunks, 1):
        slide_more(prs, cfg, it, ch, k, len(chunks), n + k)
    return 1 + len(chunks)


def slide_adjust(prs, cfg, it, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, 0.7, 0.35, 16, 1.3, [[("建議調整", {"size": 28, "bold": True, "color": C_RED}),
                                (f'　{it.get("no", "")}', {"size": 18, "color": C_GREY})]])
    line(s, 0.7, 1.75, 32.9, 1.75, C_LINE, 1)
    paras = [[("調整：", {"bold": True, "color": C_RED})]] + \
            [[(l, {"color": C_RED})] for l in str(it.get("adjust", "")).split("\n")]
    h = est_h(it.get("adjust", ""), 21) + 0.8
    tb(s, 0.9, 2.1, 12.0, h, paras)
    if it.get("replies"):
        rp = [[("回覆歷程：", {"bold": True, "color": C_GREEN, "size": 12})]] + \
             [[(l, {"color": C_GREEN, "size": 12})] for l in str(it["replies"]).split("\n")]
        tb(s, 0.9, 2.1 + h + 0.3, 12.0, 8, rp, size=12)
    line(s, 13.3, 2.1, 13.3, 17.9, C_LINE, 0.75)
    image_grid(s, it.get("adjust_images", []), 13.7, 2.1, 19.2, 15.9)
    footer(s, cfg, n)


def generate(cfg, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    items = cfg["items"]
    slide_cover(prs, cfg)
    n = 2
    if cfg.get("show_stats", True):
        slide_stats(prs, cfg, items, n); n += 1
    if cfg.get("show_index", True):
        n += slide_index(prs, cfg, items, n)
    groups = OrderedDict()
    for it in items:
        groups.setdefault(it.get("group", "") or "", []).append(it)
    for g, its in groups.items():
        if g:
            slide_big(prs, cfg, g, f"共 {len(its)} 題", n); n += 1
        for it in its:
            n += slide_question(prs, cfg, it, n)
            if it.get("adjust") or it.get("adjust_images"):
                slide_adjust(prs, cfg, it, n); n += 1
        if g:
            slide_big(prs, cfg, "END", "", n); n += 1
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return out_path


def filter_items(cfg, only=None, track=None, group=None):
    its = cfg["items"]
    if only:
        want = {x.strip().upper() for x in only.replace("，", ",").split(",") if x.strip()}
        its = [i for i in its if str(i.get("tid", "")).upper() in want or
               (str(i.get("no", "")).upper() in want and not str(i.get("tid", "")))]
    if track:
        ok = {x.strip() for x in track.replace("，", ",").split(",")}
        its = [i for i in its if i.get("track", "") in ok]
    if group:
        its = [i for i in its if i.get("group", "") == group]
    cfg["items"] = its
    return cfg


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="釋疑簡報產生器")
    ap.add_argument("src", help="問題清單 Excel 或 JSON")
    ap.add_argument("out", nargs="?", help="輸出 pptx（預設用 設定 頁的輸出檔名）")
    ap.add_argument("--only", help="只做這些追蹤編號，例：R158,R159")
    ap.add_argument("--track", help="只做這些追蹤狀態，例：預計提出 或 待回覆,已回覆待確認")
    ap.add_argument("--group", help="只做某一群組")
    ap.add_argument("--title", help="覆寫簡報標題")
    ap.add_argument("--date", help="覆寫日期")
    a = ap.parse_args()
    cfg = load_input(a.src)
    filter_items(cfg, a.only, a.track, a.group)
    if a.title:
        cfg["title"] = a.title
    if a.date:
        cfg["date"] = a.date
    if not cfg["items"]:
        sys.exit("沒有符合條件的題目")
    out = a.out or (cfg.get("output") or "釋疑簡報.pptx")
    if not os.path.isabs(out):
        out = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(a.src)), out))
    print(generate(cfg, out))
