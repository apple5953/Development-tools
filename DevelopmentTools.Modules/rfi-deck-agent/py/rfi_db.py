# -*- coding: utf-8 -*-
"""釋疑問題庫（<專案>釋疑問題庫.xlsx）管理工具

  python rfi_db.py import                     收件匣（Revit 釋疑擷取）→ 問題庫，自動給 Q 編號、複製圖片
  python rfi_db.py raise R158 R159 --date 2026-10-28
                                              標記「已在雙週會提出」：寫提出日期、回覆期限(+14天)、追蹤狀態=待回覆
  python rfi_db.py reply R158 "事務所回覆內容" [--date 2026-11-11]
                                              記錄事務所回覆 → 追蹤狀態=已回覆待確認（同時附加到回覆歷程）
  python rfi_db.py close R158 "定案結論" [--date 2026-11-11]
                                              雙方定案 → 追蹤狀態=已定案、是否解決=是
  python rfi_db.py status [--date 2026-10-28] 列出各追蹤狀態數量、逾期未回覆、本期定案

共用選項：--db 指定問題庫路徑（預設 03_資料庫/釋疑問題庫/ 裡第一個 *釋疑問題庫.xlsx；沒有就建立 釋疑問題庫.xlsx）
每次寫入前自動備份到 釋疑問題庫/_備份/。
"""
import argparse
import datetime as dt
import glob
import json
import os
import re
import shutil
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from copy import copy

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))


def default_db(root):
    """工作資料夾 → 問題庫路徑：沿用既有的 *釋疑問題庫.xlsx，沒有就用 釋疑問題庫.xlsx（第一次匯入時自動建立）"""
    d = os.path.join(root, "03_資料庫", "釋疑問題庫")
    try:
        for f in sorted(os.listdir(d)):
            if f.endswith("釋疑問題庫.xlsx") and not f.startswith("~$"):
                return os.path.join(d, f)
    except OSError:
        pass
    return os.path.join(d, "釋疑問題庫.xlsx")


DB_DEFAULT = default_db(ROOT)

# 新問題庫的欄位（與既有問題庫相同；追蹤欄位由 ensure_track_columns 補上）
BASE_COLS = ["編號", "群組", "類別", "狀態", "RFI編號", "是否解決", "圖說編號名稱", "位置", "問題", "建議", "回覆歷程",
             "圖片1", "圖說1", "圖片2", "圖說2", "圖片3", "圖說3", "圖片4", "圖說4",
             "調整說明", "調整圖片1", "調整圖說1", "調整圖片2", "調整圖說2", "備註", "對應圖說頁(自動)"]


def create_db(path):
    """建立空白問題庫（新同事第一次產生簡報時）"""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "問題清單"
    ws.append(BASE_COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5C8A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, name in enumerate(BASE_COLS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 30 if name in ("問題", "建議", "回覆歷程") else 12
    h = wb.create_sheet("欄位說明")
    h.append(["欄位", "說明"])
    wb.save(path)

TRACK_COLS = ["追蹤編號", "提出日期", "回覆期限", "事務所回覆", "定案結論", "定案日期", "追蹤狀態", "來源", "Revit視圖", "元件ID",
              "其他圖片", "其他圖說"]
TRACK_HELP = {
    "追蹤編號": "全庫唯一的識別碼 R001、R002…（編號是各群組內的序號，可能重複；追蹤、搜尋請用追蹤編號）",
    "提出日期": "在雙週會正式提出的日期（rfi_db.py raise 自動填）",
    "回覆期限": "提出日期 + 14 天（下一次雙週會）",
    "事務所回覆": "建築師事務所的回覆（rfi_db.py reply）",
    "定案結論": "雙方達成共識的最終答案（rfi_db.py close）",
    "定案日期": "定案的會議日期",
    "追蹤狀態": "預計提出 → 待回覆 → 已回覆待確認 → 已定案（另有：不提出）",
    "來源": "Revit擷取 + 收件匣資料夾名稱，或人工輸入",
    "Revit視圖": "擷取時的 Revit 視圖名稱",
    "元件ID": "擷取時選取的 Revit 元件 ID（可用來回到模型定位）",
    "其他圖片": "第 5 張以後的圖片路徑，用 ; 分隔（簡報會自動放到續頁）",
    "其他圖說": "其他圖片對應的圖說，用 ; 分隔",
}
TRACK_STATES = ["預計提出", "待回覆", "已回覆待確認", "已定案", "不提出"]
REPLY_DAYS = 14


# ------------------------------------------------------------------ 基本
def today():
    return dt.date.today()


def parse_date(s):
    if s is None or s == "":
        return None
    if isinstance(s, dt.datetime):
        return s.date()
    if isinstance(s, dt.date):
        return s
    s = str(s).strip().replace(".", "/").replace("-", "/")
    m = re.match(r"^(\d{2,4})/(\d{1,2})/(\d{1,2})", s)
    if not m:
        raise ValueError(f"看不懂的日期：{s}")
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 1000:  # 民國年
        y += 1911
    return dt.date(y, mo, d)


def fmt(d):
    return d.strftime("%Y/%m/%d") if d else ""


class DB:
    def __init__(self, path):
        if not os.path.exists(path):
            create_db(path)
        self.path = os.path.abspath(path)
        self.dir = os.path.dirname(self.path)
        self.wb = openpyxl.load_workbook(self.path)
        self.ws = self.wb["問題清單"] if "問題清單" in self.wb.sheetnames else self.wb.worksheets[0]
        self.head = [("" if c.value is None else str(c.value).strip()) for c in self.ws[1]]
        while self.head and self.head[-1] == "":
            self.head.pop()

    def col(self, name):
        return self.head.index(name) + 1

    def ensure_track_columns(self):
        added = []
        ref = self.ws.cell(1, 1)
        for name in TRACK_COLS:
            if name in self.head:
                continue
            c = len(self.head) + 1
            cell = self.ws.cell(1, c, name)
            cell.font, cell.fill, cell.alignment, cell.border = copy(ref.font), copy(ref.fill), copy(ref.alignment), copy(ref.border)
            letter = cell.column_letter
            self.ws.column_dimensions[letter].width = 30 if name in ("事務所回覆", "定案結論") else 14
            self.head.append(name)
            added.append(name)
        if added:
            # 舊資料補追蹤編號、推算追蹤狀態
            k = self.next_tid()
            for r in self.rows():
                if not self.get(r, "追蹤編號"):
                    self.set(r, "追蹤編號", f"R{k:03d}")
                    k += 1
            for r in self.rows():
                if self.get(r, "追蹤狀態"):
                    continue
                self.set(r, "追蹤狀態", derive_state(self.get(r, "狀態"), self.get(r, "是否解決"),
                                                  self.get(r, "回覆歷程"), "", ""))
            if "欄位說明" in self.wb.sheetnames:
                h = self.wb["欄位說明"]
                have = {h.cell(i, 1).value for i in range(1, h.max_row + 1)}
                for name in added:
                    if name not in have:
                        h.append([name, TRACK_HELP[name]])
        return added

    def rows(self):
        for r in range(2, self.ws.max_row + 1):
            if self.ws.cell(r, 1).value or self.ws.cell(r, self.col("問題")).value:
                yield r

    def get(self, r, name):
        if name not in self.head:
            return ""
        v = self.ws.cell(r, self.col(name)).value
        if isinstance(v, (dt.datetime, dt.date)):
            return fmt(parse_date(v))
        return "" if v is None else str(v).strip()

    def set(self, r, name, value):
        cell = self.ws.cell(r, self.col(name), value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    def find(self, key):
        """用追蹤編號（R012）找；也接受編號（Q005），但編號在不同群組會重複，重複時要求改用追蹤編號"""
        k = str(key).strip().upper()
        m = re.match(r"^R0*(\d+)$", k)
        if m:
            tid = f"R{int(m.group(1)):03d}"
            for r in self.rows():
                if self.get(r, "追蹤編號").upper() == tid:
                    return r
            raise SystemExit(f"問題庫找不到 {tid}")
        no = norm_no(k)
        hits = [r for r in self.rows() if norm_no(self.get(r, "編號")) == no]
        if len(hits) == 1:
            return hits[0]
        if not hits:
            raise SystemExit(f"問題庫找不到 {no}")
        lst = "、".join(f'{self.get(r, "追蹤編號") or "列" + str(r)}（{self.get(r, "群組")}）' for r in hits)
        raise SystemExit(f"{no} 在多個群組都有：{lst}\n請改用追蹤編號（R開頭）。")

    def label(self, r):
        return f'{self.get(r, "追蹤編號")} {self.get(r, "群組")} {self.get(r, "編號")}'.strip()

    def next_tid(self):
        mx = 0
        for r in self.rows():
            m = re.match(r"^R(\d+)", self.get(r, "追蹤編號"), re.I)
            if m:
                mx = max(mx, int(m.group(1)))
        return mx + 1

    def next_no(self, group):
        """群組內下一個 Q 編號"""
        mx = 0
        for r in self.rows():
            if self.get(r, "群組") != group:
                continue
            m = re.match(r"^Q(\d+)", self.get(r, "編號"), re.I)
            if m:
                mx = max(mx, int(m.group(1)))
        return mx + 1

    def append_row(self, rec):
        r = self.ws.max_row + 1
        # 找最後一列有資料的列，避免中間空列
        while r > 2 and not any(self.ws.cell(r - 1, c).value for c in range(1, len(self.head) + 1)):
            r -= 1
        tmpl = r - 1
        for name, val in rec.items():
            if name not in self.head:
                continue
            c = self.col(name)
            cell = self.ws.cell(r, c, val)
            src = self.ws.cell(tmpl, c)
            if src.has_style:
                cell.font, cell.border, cell.fill = copy(src.font), copy(src.border), copy(src.fill)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        return r

    def backup(self):
        bdir = os.path.join(self.dir, "_備份")
        os.makedirs(bdir, exist_ok=True)
        stem = os.path.splitext(os.path.basename(self.path))[0]
        dst = os.path.join(bdir, f"{stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(self.path, dst)
        # 只保留最近 30 份
        olds = sorted(glob.glob(os.path.join(bdir, stem + "_*.xlsx")))
        for f in olds[:-30]:
            try:
                os.remove(f)
            except OSError:
                pass
        return dst

    def save(self):
        self.backup()
        try:
            self.wb.save(self.path)
        except PermissionError:
            raise SystemExit("問題庫正在 Excel 開啟中，請先關閉 Excel 再執行。")


def norm_no(s):
    s = str(s or "").strip().upper()
    m = re.match(r"^Q?0*(\d+)$", s)
    return f"Q{int(m.group(1)):03d}" if m else s


def derive_state(status, solved, replies, office_reply, decision):
    status = (status or "").strip()
    if decision or str(solved).strip() in ("是", "已解決", "Y", "y", "1", "True"):
        return "已定案"
    if status in ("不提出", "暫不提出"):
        return "不提出"
    if office_reply:
        return "已回覆待確認"
    if status == "已提出":
        return "已回覆待確認" if replies else "待回覆"
    return "預計提出"


# ------------------------------------------------------------------ import
def inbox_dir(db):
    return os.path.join(db.dir, "收件匣")


def cmd_import(db, args):
    inbox = inbox_dir(db)
    os.makedirs(inbox, exist_ok=True)
    pend = sorted(d for d in os.listdir(inbox)
                  if not d.startswith("_") and os.path.isfile(os.path.join(inbox, d, "item.json")))
    only = getattr(args, "only", None)
    if only:
        want = {os.path.basename(os.path.normpath(x)) for x in only}
        pend = [d for d in pend if d in want]
    if not pend:
        print("收件匣沒有待匯入的項目。")
        return []
    db.ensure_track_columns()
    done_dir = os.path.join(inbox, "_已匯入")
    img_dir = os.path.join(db.dir, "圖片")
    os.makedirs(done_dir, exist_ok=True)
    os.makedirs(img_dir, exist_ok=True)
    seen = {db.get(r, "來源") for r in db.rows()}
    group = args.group or "預計提出"
    n = db.next_no(group)
    t = db.next_tid()
    imported = []
    for name in pend:
        folder = os.path.join(inbox, name)
        src_tag = f"Revit擷取 {name}"
        if src_tag in seen:
            print(f"  略過（已匯入過）：{name}")
            shutil.move(folder, unique(os.path.join(done_dir, name)))
            continue
        it = json.load(open(os.path.join(folder, "item.json"), encoding="utf-8-sig"))
        no, tid = f"Q{n:03d}", f"R{t:03d}"
        n += 1
        t += 1
        shots = it.get("Shots") or []
        rec = {
            "編號": no,
            "追蹤編號": tid,
            "群組": group,
            "類別": (it.get("Category") or "")[:1].upper(),
            "狀態": it.get("Status") or "預計提出",
            "圖說編號名稱": it.get("Drawing", ""),
            "位置": " ".join(x for x in [it.get("Building", ""), it.get("Location", "")] if x),
            "問題": it.get("Question", ""),
            "建議": it.get("Suggest", ""),
            "追蹤狀態": "預計提出",
            "來源": src_tag,
            "Revit視圖": "、".join(dict.fromkeys(s.get("View", "") for s in shots if s.get("View"))),
            "元件ID": ", ".join(it.get("ElementIds") or []),
        }
        if rec["狀態"] in ("不提出", "暫不提出"):
            rec["追蹤狀態"] = "不提出"
        notes = []
        for k, s in enumerate(shots, 1):
            src = os.path.join(folder, s.get("File", ""))
            if not os.path.isfile(src):
                # Inbox.Commit 會改名為 img_NN.png
                cand = os.path.join(folder, f"img_{k:02d}.png")
                src = cand if os.path.isfile(cand) else None
            if not src:
                continue
            dst_name = f"{tid}_{k}.png"
            shutil.copy2(src, os.path.join(img_dir, dst_name))
            if k <= 4:
                rec[f"圖片{k}"] = f"圖片/{dst_name}"
                rec[f"圖說{k}"] = s.get("Caption", "")
            else:
                rec["其他圖片"] = (rec.get("其他圖片", "") + ";" if rec.get("其他圖片") else "") + f"圖片/{dst_name}"
                rec["其他圖說"] = (rec.get("其他圖說", "") + ";" if rec.get("其他圖說") else "") + (s.get("Caption", "") or "").replace(";", "；")
        els = it.get("Elements") or []
        if els:
            notes.append("元件：" + "；".join(els[:3]) + (f" 等 {len(els)} 個" if len(els) > 3 else ""))
        if it.get("Author"):
            notes.append(f"擷取：{it['Author']} {it.get('CreatedAt', '')}".strip())
        rec["備註"] = "\n".join(notes)
        db.append_row(rec)
        shutil.move(folder, unique(os.path.join(done_dir, name)))
        imported.append(tid)
        print(f"  ✔ {tid}（{group} {no}）  {rec['位置']}  {rec['問題'][:30]}")
    if imported:
        db.save()
        log = os.path.join(db.dir, "收件匣", "_已匯入", "匯入紀錄.txt")
        with open(log, "a", encoding="utf-8") as f:
            f.write(f"{dt.datetime.now():%Y-%m-%d %H:%M}  {', '.join(imported)}\n")
    print(f"匯入 {len(imported)} 筆：{', '.join(imported) if imported else '—'}")
    if args.json_out:
        json.dump({"imported": imported}, open(args.json_out, "w", encoding="utf-8"), ensure_ascii=False)
    return imported


def unique(path):
    if not os.path.exists(path):
        return path
    i = 2
    while os.path.exists(f"{path}_{i}"):
        i += 1
    return f"{path}_{i}"


# ------------------------------------------------------------------ tracking
def cmd_raise(db, args):
    db.ensure_track_columns()
    d = parse_date(args.date) or today()
    for no in args.nos:
        r = db.find(no)
        db.set(r, "提出日期", fmt(d))
        db.set(r, "回覆期限", fmt(d + dt.timedelta(days=REPLY_DAYS)))
        db.set(r, "狀態", "已提出")
        if args.rfi:
            db.set(r, "RFI編號", args.rfi)
        if db.get(r, "追蹤狀態") in ("", "預計提出"):
            db.set(r, "追蹤狀態", "待回覆")
        print(f"  {db.label(r)} 已提出 {fmt(d)}，回覆期限 {fmt(d + dt.timedelta(days=REPLY_DAYS))}")
    db.save()


def _append_history(db, r, d, text):
    old = db.get(r, "回覆歷程")
    line = f"{d.month}/{d.day}-{text}"
    db.set(r, "回覆歷程", (old + "\n" if old else "") + line)


def cmd_reply(db, args):
    db.ensure_track_columns()
    d = parse_date(args.date) or today()
    r = db.find(args.no)
    db.set(r, "事務所回覆", args.text)
    _append_history(db, r, d, args.text)
    if db.get(r, "追蹤狀態") != "已定案":
        db.set(r, "追蹤狀態", "已回覆待確認")
    db.save()
    print(f"  {db.label(r)} 已記錄事務所回覆（{fmt(d)}）")


def cmd_close(db, args):
    db.ensure_track_columns()
    d = parse_date(args.date) or today()
    r = db.find(args.no)
    db.set(r, "定案結論", args.text)
    db.set(r, "定案日期", fmt(d))
    db.set(r, "是否解決", "是")
    db.set(r, "追蹤狀態", "已定案")
    _append_history(db, r, d, "定案：" + args.text)
    db.save()
    print(f"  {db.label(r)} 已定案（{fmt(d)}）")


def summarize(db, ref_date=None, since=None):
    """回傳雙週會用的統計。ref_date=本次會議日；since=上次會議日（本期定案用）"""
    ref = parse_date(ref_date) or today()
    since = parse_date(since) or (ref - dt.timedelta(days=REPLY_DAYS))
    out = {"date": fmt(ref), "since": fmt(since), "counts": {k: 0 for k in TRACK_STATES},
           "overdue": [], "open": [], "replied": [], "closed_this_period": [], "new_this_period": []}
    for r in db.rows():
        st = db.get(r, "追蹤狀態") or derive_state(db.get(r, "狀態"), db.get(r, "是否解決"), db.get(r, "回覆歷程"),
                                                db.get(r, "事務所回覆"), db.get(r, "定案結論"))
        out["counts"][st] = out["counts"].get(st, 0) + 1
        item = {k: db.get(r, k) for k in ["追蹤編號", "編號", "群組", "類別", "狀態", "RFI編號", "圖說編號名稱", "位置", "問題",
                                          "建議", "提出日期", "回覆期限", "事務所回覆", "定案結論", "定案日期"]}
        item["追蹤狀態"] = st
        due = parse_date(item["回覆期限"]) if item["回覆期限"] else None
        if st == "待回覆" and due and due < ref:
            item["逾期天數"] = (ref - due).days
            out["overdue"].append(item)
        if st in ("待回覆", "已回覆待確認"):
            out["open"].append(item)
        if st == "已回覆待確認":
            out["replied"].append(item)
        fd = parse_date(item["定案日期"]) if item["定案日期"] else None
        if st == "已定案" and fd and since < fd <= ref:
            out["closed_this_period"].append(item)
        pd_ = parse_date(item["提出日期"]) if item["提出日期"] else None
        if pd_ and since < pd_ <= ref:
            out["new_this_period"].append(item)
    # 未定案排序：逾期 → 已回覆待確認 → 有期限（近的先）→ 舊資料（無提出日期）
    def key(i):
        due = parse_date(i["回覆期限"]) if i["回覆期限"] else None
        od = i.get("逾期天數", 0)
        return (0 if od else 1 if i["追蹤狀態"] == "已回覆待確認" else 2 if due else 3,
                -od, due or dt.date(9999, 1, 1), i["追蹤編號"])
    out["open"].sort(key=key)
    out["overdue"].sort(key=lambda i: -i["逾期天數"])
    return out


def cmd_status(db, args):
    s = summarize(db, args.date, args.since)
    print(f"釋疑追蹤（{s['since']} ~ {s['date']}）")
    for k in TRACK_STATES:
        print(f"  {k:<8} {s['counts'].get(k, 0):>4}")
    print(f"  逾期未回覆 {len(s['overdue'])}：" + "、".join(f"{i['追蹤編號']}({i['逾期天數']}天)" for i in s["overdue"]))
    print(f"  本期定案 {len(s['closed_this_period'])}：" + "、".join(i["追蹤編號"] for i in s["closed_this_period"]))
    print(f"  本期提出 {len(s['new_this_period'])}：" + "、".join(i["追蹤編號"] for i in s["new_this_period"]))
    if args.json_out:
        json.dump(s, open(args.json_out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


# ------------------------------------------------------------------ main
def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description="釋疑問題庫管理")
    p.add_argument("--db", default=DB_DEFAULT)
    sub = p.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("import"); a.add_argument("--group", default=""); a.add_argument("--json-out")
    a = sub.add_parser("raise"); a.add_argument("nos", nargs="+"); a.add_argument("--date"); a.add_argument("--rfi")
    a = sub.add_parser("reply"); a.add_argument("no"); a.add_argument("text"); a.add_argument("--date")
    a = sub.add_parser("close"); a.add_argument("no"); a.add_argument("text"); a.add_argument("--date")
    a = sub.add_parser("status"); a.add_argument("--date"); a.add_argument("--since"); a.add_argument("--json-out")
    a = sub.add_parser("init")  # 只補追蹤欄位
    args = p.parse_args(argv)
    db = DB(args.db)
    if args.cmd == "import":
        cmd_import(db, args)
    elif args.cmd == "raise":
        cmd_raise(db, args)
    elif args.cmd == "reply":
        cmd_reply(db, args)
    elif args.cmd == "close":
        cmd_close(db, args)
    elif args.cmd == "status":
        cmd_status(db, args)
    elif args.cmd == "init":
        added = db.ensure_track_columns()
        if added:
            db.save()
        print("新增欄位：" + ("、".join(added) if added else "（已存在）"))


if __name__ == "__main__":
    main()
