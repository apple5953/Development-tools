# -*- coding: utf-8 -*-
"""
共用 PowerPoint 工具：複製投影片、改字保留格式、圖片等比放入框、表格增刪列。
所有簡報產生器 (gen_*.py) 都 import 這支。
"""
import copy
import re
from pptx import Presentation
from pptx.util import Emu, Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# ---------------------------------------------------------------- slides
def clone_slide(prs, src):
    """把 src 投影片完整複製到簡報最後 (含圖片關聯)。"""
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):  # 移除版面自帶的預留位置
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rel in src.part.rels.values():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            new = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new = dst.part.relate_to(rel._target, rel.reltype)
        rid_map[rel.rId] = new
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        new_el = copy.deepcopy(el)
        for node in new_el.iter():
            for att in list(node.attrib):
                if att.startswith("{%s}" % R_NS) and node.attrib[att] in rid_map:
                    node.attrib[att] = rid_map[node.attrib[att]]
        dst.shapes._spTree.append(new_el)
    # 預留位置若沒有自己的座標，補上版面配置的座標 (避免部分軟體位置跑掉)
    for shp in dst.placeholders:
        try:
            if shp._element.spPr.find("{%s}xfrm" % A_NS) is None:
                lp = shp._base_placeholder
                if lp is not None and lp.left is not None:
                    shp.left, shp.top, shp.width, shp.height = lp.left, lp.top, lp.width, lp.height
        except Exception:
            pass
    # 背景
    bg = src._element.cSld.bg
    if bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(bg))
    return dst


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sid in list(sld_ids):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            sld_ids.remove(sid)
            return


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from iter_shapes(s.shapes)


def find_shape(slide, name=None, startswith=None, contains=None):
    for s in iter_shapes(slide.shapes):
        if name and s.name == name:
            return s
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if startswith and t.startswith(startswith):
                return s
            if contains and contains in t:
                return s
    return None

# ---------------------------------------------------------------- rich text
# 行內標記：[紅]文字[/紅] 紅字強調；[重點]文字[/重點] 綠色底線(建議重點)；[粗]文字[/粗] 粗體；[藍]文字[/藍]
TAG_RE = re.compile(r"\[(紅|重點|粗|藍)\](.*?)\[/\1\]", re.S)


def parse_rich(text):
    """'一般[紅]紅字[/紅]一般' -> [('一般',None),('紅字','紅'),('一般',None)]"""
    out, pos = [], 0
    for m in TAG_RE.finditer(text or ""):
        if m.start() > pos:
            out.append((text[pos:m.start()], None))
        out.append((m.group(2), m.group(1)))
        pos = m.end()
    if pos < len(text or ""):
        out.append((text[pos:], None))
    return out or [("", None)]


def _style_run(run, style):
    if style == "紅":
        run.font.color.rgb = RGBColor(0xFF, 0, 0)
    elif style == "重點":
        run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_6
        run.font.underline = True
    elif style == "粗":
        run.font.bold = True
    elif style == "藍":
        run.font.color.rgb = RGBColor(0, 0, 0xFF)


def _clear_runs(p):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for br in p._p.findall("{%s}br" % A_NS):
        p._p.remove(br)


def set_paragraphs(shape, paras, keep_label_runs=0, rich=True, base="label"):
    """
    依範本段落格式改寫文字框。
    paras: list，每個元素為一段文字 (可含行內標記)。
    範本的第 i 段格式套用到第 i 段；段數不足時沿用最後一段格式。
    keep_label_runs: 每段保留前 n 個 run (例如『疑義澄清問題說明：』標籤)，文字接在後面。
    """
    tf = shape.text_frame
    tpl_ps = list(tf.paragraphs)
    tpl_p_xml = [copy.deepcopy(p._p) for p in tpl_ps]
    body = tf._txBody
    for p in tpl_ps:
        body.remove(p._p)
    for i, text in enumerate(paras):
        src = tpl_p_xml[min(i, len(tpl_p_xml) - 1)]
        new_p = copy.deepcopy(src)
        body.append(new_p)
        p = tf.paragraphs[-1]
        runs = list(p.runs)
        base_rpr = None
        body_run = runs[keep_label_runs] if len(runs) > keep_label_runs else (runs[-1] if runs else None)
        if base == "plain":  # 以第一個非粗體 run 作為內文格式
            for rr in runs:
                rp = rr._r.find("{%s}rPr" % A_NS)
                if rp is None or rp.get("b") not in ("1", "true"):
                    body_run = rr
                    break
        if body_run is not None and body_run._r.find("{%s}rPr" % A_NS) is not None:
            base_rpr = copy.deepcopy(body_run._r.find("{%s}rPr" % A_NS))
        for r in runs[keep_label_runs:]:
            r._r.getparent().remove(r._r)
        segs = parse_rich(text) if rich else [(text, None)]
        for seg, style in segs:
            lines = seg.split("\n")
            for li, line in enumerate(lines):
                if li > 0:
                    p.add_line_break()
                r = p.add_run()
                r.text = line
                if base_rpr is not None:
                    old = r._r.find("{%s}rPr" % A_NS)
                    if old is not None:
                        r._r.remove(old)
                    nr = copy.deepcopy(base_rpr)
                    # 標記樣式不要繼承範本的底線/顏色
                    if style:
                        for k in ("u",):
                            if k in nr.attrib:
                                del nr.attrib[k]
                        for child in list(nr):
                            if child.tag.endswith("solidFill"):
                                nr.remove(child)
                    r._r.insert(0, nr)
                _style_run(r, style)


def set_text_keep_format(shape, text):
    """單段文字取代，保留第一個 run 的格式。"""
    set_paragraphs(shape, [text])

# ---------------------------------------------------------------- pictures
def add_picture_fit(slide, path, left, top, width, height, align="center"):
    """圖片等比縮放後置中放進指定框 (EMU)。回傳 picture shape。"""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    box_r = width / height
    img_r = iw / ih
    if img_r > box_r:
        w = width
        h = int(width / img_r)
    else:
        h = height
        w = int(height * img_r)
    x = left + (width - w) // 2
    y = top + (height - h) // 2 if align == "center" else top
    return slide.shapes.add_picture(path, x, y, w, h)


def add_red_box(slide, left, top, width, height, line_pt=2.25):
    """紅色框 (標示問題位置)。座標 EMU。"""
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    s.fill.background()
    s.line.color.rgb = RGBColor(0xFF, 0, 0)
    s.line.width = Pt(line_pt)
    s.shadow.inherit = False
    return s

# ---------------------------------------------------------------- tables
def table_resize_rows(table, header_rows, n_data):
    """保留前 header_rows 列，資料列以最後一列為樣板增減至 n_data 列。"""
    tbl = table._tbl
    trs = tbl.findall("{%s}tr" % A_NS)
    data = trs[header_rows:]
    proto = copy.deepcopy(data[-1] if data else trs[-1])
    for tr in data:
        tbl.remove(tr)
    for _ in range(n_data):
        tbl.append(copy.deepcopy(proto))


def set_cell(cell, text, color=None):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    rpr = None
    if runs and runs[0]._r.find("{%s}rPr" % A_NS) is not None:
        rpr = copy.deepcopy(runs[0]._r.find("{%s}rPr" % A_NS))
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _clear_runs(p)
    r = p.add_run()
    r.text = "" if text is None else str(text)
    if rpr is not None:
        r._r.insert(0, rpr)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def emu(cm):
    return int(Cm(cm))
