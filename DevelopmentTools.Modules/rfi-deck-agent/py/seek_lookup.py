# -*- coding: utf-8 -*-
"""圖說查找：用 Revit 擷取的情境（棟別、樓層、柱線、元件）在「圖說查找索引」找對應的 PDF 圖說，並裁切到問題位置

  python seek_lookup.py --item 收件匣/xxx/item.json --out 輸出資料夾 [--top 4]
  python seek_lookup.py --building 實驗棟 --level 3F --grid "LINE(4-5, B-C)" --cats 結構構架 --out out

輸出：out/seek_result.json ＋ out/cand_N.png（裁切圖，紅框＝柱線範圍、紅圈＝物件位置；有結構柱時自動校正圖說偏移）
另外回傳 history：釋疑問題庫中同棟、同樓層、柱線 ±1 格內的舊問題。
索引與 PDF 位置：預設 <AI AGENT>/03_資料庫/圖說/圖說查找索引.db 與該資料夾下放 PDF 的子資料夾；
專案設定（棟別名稱、柱線對應棟別、柱線座標檔、釋疑問題庫）：<AI AGENT>/03_資料庫/圖說/查圖設定.json。
本程式也內嵌在 Revit 外掛裡（外掛更新即更新規則），外掛會以 --root 傳入 AI AGENT 資料夾。
別台電腦 PDF 放別處 → <DEVELOPMENTTOOLS_DATA 或 %APPDATA%>/DT_AIAgent/seek.json 寫 {"pdf_dir": "...", "index": "..."}。
只讀索引，不做 OCR，所以每台電腦結果相同（前提：PDF 版本相同，程式會比對檔案大小與特徵碼）。
"""
import argparse
import json
import os
import re
import sqlite3
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = DWG = DEFAULT_INDEX = DEFAULT_PDF = ""
CFG = {}


def set_root(root):
    """AI AGENT 資料夾 → 各路徑與專案設定（查圖設定.json）"""
    global ROOT, DWG, DEFAULT_INDEX, DEFAULT_PDF, CFG, _MASTER, LV_TXT
    ROOT = os.path.normpath(root)
    DWG = os.path.join(ROOT, "03_資料庫", "圖說")
    DEFAULT_INDEX = os.path.join(DWG, "圖說查找索引.db")
    try:
        CFG = json.load(open(os.path.join(DWG, "查圖設定.json"), encoding="utf-8-sig"))
    except Exception:
        CFG = {}
    sub = CFG.get("pdf_subdir", "")
    if not sub and os.path.isdir(DWG):   # 沒設定：第一個放了 PDF 的子資料夾
        for d in sorted(os.listdir(DWG)):
            dd = os.path.join(DWG, d)
            if os.path.isdir(dd) and any(f.lower().endswith(".pdf") for f in os.listdir(dd)):
                sub = d
                break
    DEFAULT_PDF = os.path.join(DWG, sub) if sub else DWG
    _MASTER = None
    names = "|".join(re.escape(b) for b in CFG.get("buildings", [])) or "(?!x)x"
    LV_TXT = re.compile(r"(" + names + r")?(屋突)?(地下)?([一二三四五六七八九十]+)[層樓]")


def first_file(folder, suffix):
    try:
        return next((os.path.join(folder, f) for f in sorted(os.listdir(folder)) if f.endswith(suffix)), "")
    except OSError:
        return ""


CAT_DISC = [  # Revit 類別關鍵字 → 圖號前綴（專業）
    (("結構", "基礎", "鋼筋"), ["S"]),
    (("風管", "空調", "機械設備", "風口"), ["AC"]),
    (("消防", "撒水", "灑水"), ["F"]),
    (("電纜", "電氣", "燈", "照明", "線槽", "開關", "插座", "管道"), ["E", "T"]),
    (("管", "衛浴", "給水", "排水", "衛工", "管件", "管配件"), ["P", "W"]),
    (("牆", "門", "窗", "樓板", "天花", "樓梯", "欄杆", "屋頂", "房間", "帷幕", "家具"), ["A"]),
]


def settings():
    # 外掛啟動時設定 DEVELOPMENTTOOLS_DATA（非系統碟）；找不到時退回舊的 %APPDATA%。
    p = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "DT_AIAgent", "seek.json")
    data = os.environ.get("DEVELOPMENTTOOLS_DATA")
    if data and os.path.exists(os.path.join(data, "DT_AIAgent", "seek.json")):
        p = os.path.join(data, "DT_AIAgent", "seek.json")
    try:
        return json.load(open(p, encoding="utf-8"))
    except Exception:
        return {}


def norm_level(s):
    """Revit 樓層名 → 索引格式：B3F / 3F / RF / PRF"""
    t = (s or "").upper().replace("Ｆ", "F")
    if re.search(r"屋突|PR", t):
        return "PRF"
    m = re.search(r"B\s*(\d+)", t)
    if m:
        return f"B{int(m.group(1))}F"
    if re.search(r"\bR\d*F|RF|屋頂|屋面|\bR\d", t):
        return "RF"
    m = re.search(r"(\d+)\s*F", t) or re.search(r"(\d+)", t)
    if m:
        return f"{int(m.group(1))}F"
    m = re.search(r"地下([一二三四五])", s or "")
    if m:
        return f'B{"一二三四五".index(m.group(1)) + 1}F'
    return ""


def parse_grid(s):
    """'LINE(8-9, G'-H)' → [['8','9'], ["G'",'H']]"""
    m = re.search(r"\((.*)\)", s or "")
    body = m.group(1) if m else (s or "")
    out = []
    for part in body.split(","):
        labs = [x.strip() for x in re.split(r"[-~～]", part.strip()) if x.strip()]
        if labs:
            out.append(labs)
    return out


CAT_WORDS = [  # Revit 類別 → 圖名關鍵字加分
    (("風管", "風口"), "風管"), (("消防", "撒水", "灑水"), "撒水"), (("電纜", "線槽"), "電"),
    (("給水",), "給水"), (("排水", "衛浴"), "排水"), (("結構",), "結構"), (("天花",), "天花"),
]
CN = "一二三四五六七八九十"


def level_cn(lv):
    """B2F → 地下二層；3F → 三層"""
    m = re.match(r"^(B)?(\d+)F$", lv or "")
    if not m:
        return ""
    n = int(m.group(2))
    t = CN[n - 1] if n <= 10 else "十" + CN[n - 11]
    return ("地下" if m.group(1) else "") + t + "層"


def want_prefixes(cats):
    pre = []
    for c in cats:
        for keys, ps in CAT_DISC:
            if any(k in c for k in keys):
                pre += ps
                break
    return list(dict.fromkeys(pre))


MARK = re.compile(r"\b[A-Z]{1,4}\d{1,3}[A-Z]?\b")


def guess_building(gparts):
    """查圖設定.json 的 building_by_grid：[[起, 迄, "棟別"], …]（數字柱線範圍對應棟別）"""
    for part in gparts:
        for lab in part:
            if lab.isdigit():
                for lo, hi, b in CFG.get("building_by_grid", []):
                    if lo <= int(lab) <= hi:
                        return b
    return ""


def find(con, building, level, grid, cats, types, keywords, top, refs=None, labels=None):
    pres = want_prefixes(cats) or ["A", "S"]
    if not building:
        building = guess_building(parse_grid(grid))
    lv = norm_level(level)
    gparts = parse_grid(grid)
    want_labels = {l for part in gparts for l in part}
    marks = sorted({m for t in types for m in MARK.findall((t or "").upper())})
    rows = con.execute("SELECT id,file,page,w,h,drawing_no,title,prefix,building,levels,kind,grid_v,grid_h FROM pages").fetchall()
    title_of = {r[0]: r[6] for r in rows}
    wrong_level = set()
    cands = []
    for r in rows:
        pid, f, pg, w, h, no, title, prefix, bld, levels, kind, gv, gh = r
        score, why = 0, []
        lvls = (levels or "").split(",")
        tl = _text_levels(title)
        if tl:
            # 圖名寫的樓層為準：索引沒抓到「一樓」、也把「屋突一層」誤當 1F
            lvls = sorted({l for _, l in tl})
        if prefix in pres:
            score += 30 if prefix == pres[0] else 22; why.append(f"專業 {prefix}")
        elif prefix in ("A",) and not cats:
            score += 10
        elif cats and prefix != "A":
            score -= 20         # 選的是結構構件，機電圖往後排（建築圖保留當第二參考）
        if lv and lv in lvls:
            score += 40; why.append(f"樓層 {lv}")
        elif lv and lvls and lvls != [""]:
            score -= 25
        if building and building in (bld or ""):
            score += 20; why.append(building)
        elif building and bld:
            score -= 15
        if kind == "平面圖":
            score += 15
            # 主平面圖（「實驗棟一樓平面圖(1)」）優先於插座、指標、地坪等分項平面圖
            if re.search(r"[層樓]平面圖", title or ""):
                score += 10; why.append("主平面圖")
        # 棟別＋樓層要連在一起（例：「實驗棟地下二層及行政棟地下一層」→ 實驗棟是 B2F）
        lc = level_cn(lv)
        # 圖名寫的樓層要對（「屋突一層」不是 1F、「筏基層」不是 B1F）
        if lv and tl and not any(l == lv and (b is None or not building or b == building) for b, l in tl):
            score -= 40
            wrong_level.add(pid)
        elif building and lc:
            if building + lc in (title or "") or (building, lv) in tl:
                score += 20; why.append(building + lc)
            elif building in (title or "") and lc in (title or ""):
                score -= 10
        for keys, word in CAT_WORDS:
            if any(k in c for c in cats for k in keys) and word in (title or ""):
                score += 10; why.append("圖名含「" + word + "」")
                break
        gv, gh = json.loads(gv or "[]"), json.loads(gh or "[]")
        labs = {g[0] for g in gv + gh}
        hit = want_labels & labs
        if want_labels and hit:
            score += 5 * len(hit); why.append("柱線 " + "、".join(sorted(hit)))
        cands.append([score, pid, why])
    # 構件編號（例：G3、SG1）→ 加分詳圖/大樣
    if marks:
        q = ",".join("?" * len(marks))
        for pid, t in con.execute(f"SELECT page_id, text FROM marks WHERE text IN ({q})", marks):
            for c in cands:
                if c[1] == pid:
                    c[0] += 4           # 元件類型名稱常與圖說不同，只當參考
                    if "構件 " + t not in c[2]:
                        c[2].append("構件 " + t)
    # 柱編號（模型備註欄）：同一張圖上找到 ≥3 支附近柱的編號 → 這張圖就畫到這一區
    if refs:
        q = ",".join("?" * len(refs))
        hits = {}
        for pid, t in con.execute(f"SELECT page_id, UPPER(text) FROM marks WHERE UPPER(text) IN ({q})", list(refs)):
            hits.setdefault(pid, set()).add(t)
        for c in cands:
            h = hits.get(c[1])
            if h and len(h) >= 3 and c[1] not in wrong_level:
                c[0] += 25
                c[2].append("柱編號 " + "、".join(sorted(h)[:4]) + ("…" if len(h) > 4 else ""))
    # 選取構件的編號（梁 G4e、J7a…；全文搜尋）
    for lab in labels or []:
        if len(lab) < 3 or lab in (refs or {}):
            continue
        try:
            ids = {r[0] for r in con.execute("SELECT rowid FROM fts WHERE fts MATCH ?", ('"' + lab + '"',))}
        except sqlite3.OperationalError:
            ids = set()
        for c in cands:
            if c[1] in ids:
                c[0] += 8
                c[2].append("編號 " + lab)
    # 關鍵字全文
    for kw in keywords:
        kw = kw.strip()
        if len(kw) < 3:
            continue
        try:
            ids = {r[0] for r in con.execute("SELECT rowid FROM fts WHERE fts MATCH ?", ('"' + kw.replace('"', "") + '"',))}
        except sqlite3.OperationalError:
            ids = set()
        for c in cands:
            if c[1] in ids:
                c[0] += 8
                c[2].append("文字 " + kw)
    cands.sort(key=lambda c: -c[0])
    # 只找結構（S）與建築（A）圖說；前幾名兩種都要有
    pre_of = {r[0]: r[7] for r in rows}
    pool = [c for c in cands if c[0] > 0 and pre_of.get(c[1]) in ("S", "A")]
    pick = pool[:top]
    for want in ("S", "A"):
        if pick and not any(pre_of[c[1]] == want for c in pick):
            extra = next((c for c in pool[top:] if pre_of[c[1]] == want), None)
            if extra:
                # 換掉同專業中分數最低的那張
                major = max(("S", "A"), key=lambda p: sum(pre_of[c[1]] == p for c in pick))
                for i in range(len(pick) - 1, -1, -1):
                    if pre_of[pick[i][1]] == major:
                        pick[i] = extra
                        break
    pick.sort(key=lambda c: -c[0])
    return pick, gparts, marks


def grid_box(gv, gh, gparts):
    """柱線範圍 → 頁面比例框 (x0,y0,x1,y1)；找不到回傳 None"""
    vmap, hmap = {g[0]: g[1] for g in gv}, {g[0]: g[1] for g in gh}
    xs, ys = [], []
    for part in gparts:
        if all(l in vmap for l in part):
            xs = [vmap[l] for l in part]
        elif all(l in hmap for l in part):
            ys = [hmap[l] for l in part]
        else:
            # 部分找到也用
            xs = xs or [vmap[l] for l in part if l in vmap]
            ys = ys or [hmap[l] for l in part if l in hmap]
    if not xs and not ys:
        return None
    def span(vals, allv):
        if not vals:
            return None
        if len(vals) == 1:
            # 單一柱線：取相鄰柱線距離的一半當寬度
            pos = sorted(v for v in allv)
            gap = min([abs(b - a) for a, b in zip(pos, pos[1:])] or [0.05])
            return vals[0] - gap / 2, vals[0] + gap / 2
        return min(vals), max(vals)
    sx = span(xs, vmap.values()) or (0.05, 0.85)
    sy = span(ys, hmap.values()) or (0.05, 0.95)
    return sx[0], sy[0], sx[1], sy[1]


def _axis_map(gv, gh, part):
    """part 的柱線在哪個方向：回傳 (axis 'x'/'y', {label: 頁面比例})；找不到回傳 (None, {})"""
    vmap, hmap = {g[0]: g[1] for g in gv}, {g[0]: g[1] for g in gh}
    if all(l in vmap for l in part):
        return "x", vmap
    if all(l in hmap for l in part):
        return "y", hmap
    return None, {}


def grid_point(gv, gh, gparts, cellpos):
    """柱線格內相對位置 (fx,fy) → 頁面比例座標 (px,py)；任一方向算不出來就回傳 None"""
    try:
        fr = [float(v) for v in (cellpos or "").split(",")]
    except ValueError:
        return None
    pt = {}
    for k, part in enumerate(gparts[:2]):
        axis, m = _axis_map(gv, gh, part)
        if not axis:
            continue
        if len(part) == 1:
            pt[axis] = m[part[0]]
        elif k < len(fr):
            a, b = m[part[0]], m[part[-1]]
            pt[axis] = a + max(0.0, min(1.0, fr[k])) * (b - a)
    if "x" in pt and "y" in pt:
        return pt["x"], pt["y"]
    return None


def filled_shapes(page):
    """頁面上所有實心圖形（柱斷面候選）：[(cx, cy, w, h)] 顯示方向頁面比例"""
    W, H = page.rect.width, page.rect.height
    mat = page.rotation_matrix
    out = []
    try:
        drawings = page.get_drawings()
    except Exception:
        return out
    for d in drawings:
        if d.get("fill") is None:
            continue
        r = d["rect"] * mat
        w, h = r.width / W, r.height / H
        if w <= 0 or h <= 0:
            continue
        out.append(((r.x0 + r.x1) / 2 / W, (r.y0 + r.y1) / 2 / H, w, h))
    return out


_MASTER = None


def grid_master():
    """Revit 柱線座標（公尺）：{'v': {'7': X}, 'h': {"B'": Y}}"""
    global _MASTER
    if _MASTER is None:
        _MASTER = {"v": {}, "h": {}}
        try:
            gm = CFG.get("grid_master", "")
            gm = os.path.join(DWG, gm) if gm else first_file(DWG, "柱線_Revit.json")
            d = json.load(open(gm, encoding="utf-8"))
            for g in d.get("grids", []):
                if abs(abs(g.get("angle", 0)) - 90) < 1:
                    _MASTER["v"][g["name"]] = g["x0"]
                else:
                    _MASTER["h"][g["name"]] = g["y0"]
        except Exception:
            pass
    return _MASTER


def grid_map(gv, gh, point):
    """用整張圖上所有已定位的柱線（不只問題那一格）做「模型座標→圖面」線性換算，算出物件點位"""
    M = grid_master()
    if not point or not (M["v"] or M["h"]):
        return None
    X, Y = point
    out = {}
    for glist, axis in ((gv, "x"), (gh, "y")):
        best = None
        for key, val in (("v", X), ("h", Y)):
            pr = [(M[key][g[0]], g[1]) for g in glist if g[0] in M[key]]
            if len(pr) < 3:
                continue
            n = len(pr)
            mx = sum(a for a, _ in pr) / n
            my = sum(b for _, b in pr) / n
            sxx = sum((a - mx) ** 2 for a, _ in pr)
            if sxx <= 0:
                continue
            k = sum((a - mx) * (b - my) for a, b in pr) / sxx
            c = my - k * mx
            res = max(abs(k * a + c - b) for a, b in pr)
            if res > 0.006:
                continue
            if best is None or n > best[0]:
                best = (n, k * val + c)
        if best:
            out[axis] = best[1]
    if "x" in out and "y" in out:
        return out["x"], out["y"]
    return None


def column_offset(page, gv, gh, gparts, shapes=None):
    """以結構柱校正：設計方 PDF 的柱線圓圈常與實際柱位錯開。
    在目標格附近的柱線交點找「實心近方形」（柱斷面），取偏移量中位數。
    回傳 (dx, dy, 命中數)（頁面比例）；找不到可信的偏移回傳 None"""
    if not gv or not gh:
        return None
    W, H = page.rect.width, page.rect.height
    xs = sorted(g[1] for g in gv)
    ys = sorted(g[1] for g in gh)
    def med(v):
        v = sorted(v) or [0.05]
        return v[len(v) // 2]
    # 用柱距中位數（輔助柱線 B' 之類距離很小，不能用最小值）
    sx = med([b - a for a, b in zip(xs, xs[1:]) if b - a > 1e-4])
    sy = med([b - a for a, b in zip(ys, ys[1:]) if b - a > 1e-4])
    # 目標附近的柱線（前後各 2 條）
    def near(glist, part):
        labs = [g[0] for g in sorted(glist, key=lambda g: g[1])]
        idx = [labs.index(l) for l in part if l in labs]
        if not idx:
            return [g[1] for g in glist]
        lo, hi = max(0, min(idx) - 2), min(len(labs) - 1, max(idx) + 2)
        pos = {g[0]: g[1] for g in glist}
        return [pos[labs[i]] for i in range(lo, hi + 1)]
    px, py = [g[1] for g in gv], [g[1] for g in gh]
    for part in gparts[:2]:
        axis, _ = _axis_map(gv, gh, part)
        if axis == "x":
            px = near(gv, part)
        elif axis == "y":
            py = near(gh, part)
    if shapes is None:
        shapes = filled_shapes(page)
    cols = []
    for cx, cy, w, h in shapes:
        if not (0.03 * sx <= w <= 0.35 * sx and 0.03 * sy <= h <= 0.35 * sy):
            continue
        if max(w / sx, h / sy) > 2.5 * min(w / sx, h / sy):
            continue
        cols.append((cx, cy, (w / sx) * (h / sy)))
    if len(cols) < 3:
        return None
    offs = []
    for x in px:
        for y in py:
            near_c = [(cx - x, cy - y, a) for cx, cy, a in cols if abs(cx - x) <= 0.45 * sx and abs(cy - y) <= 0.45 * sy]
            if not near_c:
                continue
            big = max(c[2] for c in near_c)
            near_c = [c for c in near_c if c[2] >= 0.5 * big]     # 主柱優先，略過小柱（C0a 之類）
            dx, dy, _ = min(near_c, key=lambda c: (c[0] / sx) ** 2 + (c[1] / sy) ** 2)
            offs.append((dx, dy))
    if len(offs) < 3:
        return None
    mdx = sorted(o[0] for o in offs)[len(offs) // 2]
    mdy = sorted(o[1] for o in offs)[len(offs) // 2]
    agree = [o for o in offs if abs(o[0] - mdx) <= 0.08 * sx and abs(o[1] - mdy) <= 0.08 * sy]
    if len(agree) < max(3, len(offs) * 0.5):
        return None
    dx = sum(o[0] for o in agree) / len(agree)
    dy = sum(o[1] for o in agree) / len(agree)
    if abs(dx) < 0.02 * sx and abs(dy) < 0.02 * sy:
        dx = dy = 0.0      # 本來就對齊
    return dx, dy, len(agree)


LAB_RX = re.compile(r"^([A-Za-z]{1,4}\d{0,3}[A-Za-z]?)")


def label_key(t):
    """'G5f(-40)' → 'G5F'；'C56' → 'C56'"""
    import unicodedata
    m = LAB_RX.match(unicodedata.normalize("NFKC", t or "").strip())
    return m.group(1).upper() if m else ""


def page_tokens(page, con=None, pid=None):
    """頁面文字 [(text, x0, y0, x1, y1)]（顯示方向頁面比例）。向量字直接讀；線條字的頁面改用索引裡 OCR 的構件編號"""
    import pymupdf
    W, H = page.rect.width, page.rect.height
    mat = page.rotation_matrix
    out = []
    try:
        for x0, y0, x1, y1, w, *_ in page.get_text("words"):
            r = pymupdf.Rect(x0, y0, x1, y1) * mat
            out.append((w, r.x0 / W, r.y0 / H, r.x1 / W, r.y1 / H))
    except Exception:
        pass
    if len(out) < 30 and con is not None and pid is not None:
        out += list(con.execute("SELECT text,x0,y0,x1,y1 FROM marks WHERE page_id=?", (pid,)))
    return out


def parse_refs(s):
    """'C56:12.30,45.60;C57:…' → {'C56': (12.3, 45.6)}"""
    out = {}
    for part in (s or "").split(";"):
        if ":" not in part:
            continue
        k, xy = part.split(":", 1)
        try:
            x, y = [float(v) for v in xy.split(",")]
        except ValueError:
            continue
        k = label_key(k)
        if k:
            out[k] = (x, y)
    return out


def fit_columns(refs, toks, shapes, W, H):
    """用柱編號對位：模型柱座標 (公尺) ↔ 圖上同編號的柱 → 相似轉換 w = a·z + b（z 可能要取共軛＝鏡射）。
    需 ≥3 支柱一致（誤差 ≤0.8m）。回傳 dict 或 None"""
    import itertools
    occ = {}
    for t, x0, y0, x1, y1 in toks:
        k = label_key(t)
        if k in refs:
            occ.setdefault(k, []).append((x0, y0, x1, y1))
    if len(occ) < 3:
        return None
    cand = {}
    for k, lst in occ.items():
        for x0, y0, x1, y1 in lst:
            u, v = (x0 + x1) / 2 * W, (y0 + y1) / 2 * H
            th, tw = (y1 - y0) * H, (x1 - x0) * W
            rad = max(4 * th, 1.6 * tw)
            best = None
            for cx, cy, w, h in shapes:
                pw, ph = w * W, h * H
                if not (0.4 * th <= pw <= 8 * th and 0.4 * th <= ph <= 8 * th) or max(pw, ph) > 2.5 * min(pw, ph):
                    continue
                d = ((cx * W - u) ** 2 + (cy * H - v) ** 2) ** 0.5
                if d <= rad and (best is None or d < best[0]):
                    best = (d, cx * W, cy * H)
            cand.setdefault(k, []).append(complex(best[1], best[2]) if best else complex(u, v))
    keys = list(cand)
    zs = {k: complex(*refs[k]) for k in keys}
    best = None
    for k1, k2 in itertools.combinations(keys, 2):
        if abs(zs[k1] - zs[k2]) < 3:
            continue
        for refl in (True, False):
            z1 = zs[k1].conjugate() if refl else zs[k1]
            z2 = zs[k2].conjugate() if refl else zs[k2]
            for w1 in cand[k1]:
                for w2 in cand[k2]:
                    a = (w1 - w2) / (z1 - z2)
                    if abs(a) < 1e-6:
                        continue
                    b = w1 - a * z1
                    tol = 0.8 * abs(a)
                    inl = []
                    for k in keys:
                        z = zs[k].conjugate() if refl else zs[k]
                        p = a * z + b
                        dm, wb = min((abs(p - w), w) for w in cand[k])
                        if dm <= tol:
                            inl.append((k, z, wb))
                    if best is None or len(inl) > len(best[0]):
                        best = (inl, refl)
    if not best or len(best[0]) < 3:
        return None
    inl, refl = best
    zl = [i[1] for i in inl]
    wl = [i[2] for i in inl]
    n = len(zl)
    zm, wm = sum(zl) / n, sum(wl) / n
    a = sum((w - wm) * (z - zm).conjugate() for z, w in zip(zl, wl)) / sum(abs(z - zm) ** 2 for z in zl)
    b = wm - a * zm
    resid = max(abs(a * z + b - w) for z, w in zip(zl, wl)) / abs(a)
    if resid > 1.2:
        return None
    return {"a": a, "b": b, "refl": refl, "labels": sorted(i[0] for i in inl), "resid_m": round(resid, 2),
            "pt_per_m": abs(a)}


def model_to_page(fit, xy, W, H):
    z = complex(*xy)
    if fit["refl"]:
        z = z.conjugate()
    w = fit["a"] * z + fit["b"]
    return w.real / W, w.imag / H


def nearest_label(toks, label, px, py, W, H, maxd_pt):
    """頁面上與 (px,py) 最近的同編號文字框（小梁編號會重複，只取最近的一個）"""
    best = None
    for t, x0, y0, x1, y1 in toks:
        if label_key(t) != label:
            continue
        d = ((((x0 + x1) / 2 - px) * W) ** 2 + (((y0 + y1) / 2 - py) * H) ** 2) ** 0.5
        if d <= maxd_pt and (best is None or d < best[0]):
            best = (d, (x0, y0, x1, y1))
    return best[1] if best else None


def render(pdf_path, page_no, box, out_png, max_px=1600, gv=None, gh=None, gparts=None, cellpos="",
           point=None, refs=None, labels=None, con=None, pid=None, plan=False):
    """裁切圖說並畫紅框（範圍）、紅圈（物件位置）、橘框（選取構件的編號）。回傳 info dict
    定位優先順序：① 柱編號對位（模型座標→圖面，最準）② 柱線＋格內比例（再以結構柱校正）"""
    import pymupdf
    from PIL import Image, ImageDraw
    doc = pymupdf.open(pdf_path)
    page = doc[page_no - 1]
    W, H = page.rect.width, page.rect.height
    info = {"snapped": False, "marker": None, "method": ""}
    mark = None
    hl = []                                   # 要框起來的編號文字
    if plan:
        shapes = filled_shapes(page)
        toks = page_tokens(page, con, pid) if (refs or labels) else []
        fit = fit_columns(refs, toks, shapes, W, H) if (refs and point and len(refs) >= 3) else None
        if fit:
            mark = model_to_page(fit, point, W, H)
            half = 4.5 * fit["pt_per_m"]
            box = (mark[0] - half / W, mark[1] - half / H, mark[0] + half / W, mark[1] + half / H)
            info.update(method="柱編號", fit_labels=fit["labels"], fit_resid_m=fit["resid_m"])
        elif box and gv is not None:
            off = column_offset(page, gv, gh, gparts or [], shapes)
            info["column_check"] = [round(off[0], 4), round(off[1], 4), off[2]] if off else None
            if off and (off[0] or off[1]):
                box = (box[0] + off[0], box[1] + off[1], box[2] + off[0], box[3] + off[1])
                info["snapped"] = True
                info["snap_hits"] = off[2]
            gm = grid_map(gv, gh, point)
            mark = gm or grid_point(gv, gh, gparts or [], cellpos)
            if mark and off:
                mark = (mark[0] + off[0], mark[1] + off[1])
            if gm and box:
                # 用整張柱線換算出的點位較準：紅框改以紅圈為中心（大小不變）
                hw, hh = (box[2] - box[0]) / 2, (box[3] - box[1]) / 2
                box = (mark[0] - hw, mark[1] - hh, mark[0] + hw, mark[1] + hh)
            if mark:
                info["method"] = "柱線"
        # 選取構件的編號（梁 J7a 之類會重複）：取最靠近物件位置的一個
        if labels and toks:
            ref = mark if mark else (((box[0] + box[2]) / 2, (box[1] + box[3]) / 2) if box else None)
            if ref:
                maxd = (6 * fit["pt_per_m"]) if fit else (max(box[2] - box[0], box[3] - box[1]) * max(W, H) if box else 0)
                for lab in labels:
                    hit = nearest_label(toks, lab, ref[0], ref[1], W, H, maxd)
                    if hit:
                        hl.append(hit)
                        info.setdefault("labels_found", []).append(lab)
                        if not box:
                            box = hit
        info["marker"] = mark

    def draw(view, path, px, full):
        r = pymupdf.Rect(view[0] * W, view[1] * H, view[2] * W, view[3] * H)    # 顯示方向座標
        zoom = px / max(r.width, r.height)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=r, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        if box:
            d = ImageDraw.Draw(img)
            sx, sy = img.width / (view[2] - view[0]), img.height / (view[3] - view[1])
            lw = max(3, img.width // (320 if full else 300))
            def to_img(x, y):
                return (x - view[0]) * sx, (y - view[1]) * sy
            rect = [*to_img(box[0], box[1]), *to_img(box[2], box[3])]
            d.rectangle(rect, outline=(230, 0, 0), width=lw)
            for x0, y0, x1, y1 in hl:
                a0, b0 = to_img(x0, y0)
                a1, b1 = to_img(x1, y1)
                pad = max(4, lw * 2)
                d.rectangle([a0 - pad, b0 - pad, a1 + pad, b1 + pad], outline=(255, 120, 0), width=max(2, lw - 1))
            if mark:
                cx, cy = to_img(*mark)
                side = min(rect[2] - rect[0], rect[3] - rect[1])
                rad = max(18, side * (0.2 if full else 0.12), img.width * (0.008 if full else 0))
                d.ellipse([cx - rad, cy - rad, cx + rad, cy + rad], outline=(230, 0, 0), width=lw + 1)
                d.line([cx - rad * 0.4, cy, cx + rad * 0.4, cy], fill=(230, 0, 0), width=max(2, lw - 1))
                d.line([cx, cy - rad * 0.4, cx, cy + rad * 0.4], fill=(230, 0, 0), width=max(2, lw - 1))
        img.save(path)

    # 主圖：整張圖說（含圖框）＋紅框紅圈；另存局部放大圖 _zoom.png
    view = [0, 0, 1, 1]
    draw(view, out_png, max(max_px, 2800), True)
    if box:
        bw, bh = box[2] - box[0], box[3] - box[1]
        mx, my = max(bw * 0.6, 0.06), max(bh * 0.6, 0.06)
        zv = [max(0, box[0] - mx), max(0, box[1] - my), min(1, box[2] + mx), min(1, box[3] + my)]
        zp = out_png[:-4] + "_zoom.png"
        draw(zv, zp, max_px, False)
        info["png_zoom"] = zp
    doc.close()
    info["view"] = view
    info["box"] = list(box) if box else None
    return info


# ------------------------------------------------------------ 同位置的釋疑歷史
LETTER = re.compile(r"^([A-Z])('*)$")
GRID_TOKEN = re.compile(r"([A-Z]'*|\d{1,2})\s*[-~～]\s*([A-Z]'*|\d{1,2})(?:\s*[~～]\s*([A-Z]'*|\d{1,2})\s*[-]\s*([A-Z]'*|\d{1,2}))?")


def _lab_val(l):
    """柱線標示 → 可比較的數值：數字軸取數字；字母軸 A=1、B=2，B'=2.5"""
    l = l.strip().upper()
    if l.isdigit():
        return ("n", int(l))
    m = LETTER.match(l)
    if m:
        return ("a", ord(m.group(1)) - 64 + 0.5 * len(m.group(2)))
    return None


def _text_grid(text):
    """釋疑「位置」文字 → {'n': [數字...], 'a': [字母值...]}"""
    out = {"n": [], "a": []}
    t = (text or "").upper().replace("’", "'").replace("′", "'")
    for m in GRID_TOKEN.finditer(t):
        for g in m.groups():
            if g:
                v = _lab_val(g)
                if v:
                    out[v[0]].append(v[1])
    return out


LV_TXT = None      # set_root() 依棟別名稱建立


def _cn_num(t):
    """一 → 1、十 → 10、十二 → 12、二十 → 20"""
    if "十" in t:
        a, _, b = t.partition("十")
        return (CN.index(a) + 1 if a else 1) * 10 + (CN.index(b) + 1 if b else 0)
    return CN.index(t) + 1 if t in CN else 0


def _text_levels(text):
    """文字中的樓層 → {(棟別或None, 'B2F'), ...}"""
    out = set()
    t = text or ""
    for m in re.finditer(r"(?<![A-Z0-9])(B?\d{1,2})F", t.upper()):
        out.add((None, m.group(1) + "F"))
    for m in re.finditer(r"(?<!\d)(\d{1,2})樓", t):
        out.add((None, m.group(1) + "F"))
    if "筏基" in t:
        out.add((None, "筏基"))
    for m in LV_TXT.finditer(t):
        lv = "PRF" if m.group(2) else ("B" if m.group(3) else "") + f"{_cn_num(m.group(4))}F"
        out.add((m.group(1), lv))
    if "屋頂" in t:
        out.add((None, "RF"))
    bs = CFG.get("buildings", [])
    if any(b1 + "及" + b2 in t for b1 in bs for b2 in bs if b1 != b2):   # 「A棟及B棟一層」→ 兩棟共用
        out = {(None, lv) for _, lv in out}
    return out


def find_history(building, level, gparts, limit=10):
    """在釋疑問題庫找同棟、同樓層、柱線相近（±1 格）的舊問題"""
    db = os.path.join(ROOT, CFG["rfi_db"]) if CFG.get("rfi_db") else \
        first_file(os.path.join(ROOT, "03_資料庫", "釋疑問題庫"), "釋疑問題庫.xlsx")
    if not gparts or not os.path.exists(db):
        return []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(db, read_only=True, data_only=True)
    except Exception:
        return []
    ws = wb["問題清單"] if "問題清單" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rows, [])]
    col = {h: i for i, h in enumerate(head)}
    def g(r, k):
        i = col.get(k)
        v = r[i] if i is not None and i < len(r) else None
        return "" if v is None else str(v).strip()
    want = {"n": [], "a": []}
    for part in gparts:
        for l in part:
            v = _lab_val(l)
            if v:
                want[v[0]].append(v[1])
    if not want["n"] and not want["a"]:
        wb.close()
        return []
    lv = norm_level(level)
    out = []
    for r in rows:
        pos = g(r, "位置")
        title = g(r, "圖說編號名稱")
        q = g(r, "問題")
        if not pos and not q:
            continue
        got = _text_grid(pos)
        if not got["n"] and not got["a"]:
            continue
        ok, dist = True, 0.0
        for k in ("n", "a"):
            if want[k] and got[k]:
                dk = min(abs(a - b) for a in want[k] for b in got[k])
                dist += dk
                if dk > 1.01:
                    ok = False
            elif want[k] and not got[k] and k == "n":
                ok = False          # 沒有數字柱線就無法確認是哪一區
        if not ok:
            continue
        # 棟別
        txt = pos + " " + title
        other = [b for b in CFG.get("buildings", []) if b in pos and b != building]
        if building and other and building not in pos:
            continue
        # 樓層：有寫就必須相符
        lvls = _text_levels(pos) or _text_levels(title)      # 位置有寫樓層就以位置為準
        level_note = ""
        if lv and lvls:
            if not any(l == lv and (b is None or not building or b == building) for b, l in lvls):
                continue
        elif lv:
            level_note = "（未註明樓層）"
        tid = g(r, "追蹤編號") or "-".join(re.sub(r"\s+", "", x) for x in (g(r, "群組"), g(r, "編號")) if x)
        pos = pos.splitlines()[0].strip() if pos else pos
        state = g(r, "追蹤狀態") or ("已解決" if g(r, "是否解決") == "是" else (g(r, "狀態") or "未結"))
        concl = g(r, "定案結論")
        if not concl:
            hist = [x for x in g(r, "回覆歷程").splitlines() if x.strip()]
            concl = hist[-1] if hist else ""
        summ = re.sub(r"\s+", " ", q)[:40]
        out.append({"tid": tid, "state": state, "position": pos.replace("\n", " "), "drawing": title.replace("\n", " "),
                    "summary": summ + level_note, "conclusion": concl[:80],
                    "_d": dist + (0.5 if level_note else 0) + (0.3 if state == "已解決" else 0)})
    wb.close()
    out.sort(key=lambda h: h["_d"])
    for h in out:
        h.pop("_d")
    return out[:limit]


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--item")
    ap.add_argument("--root", default="", help="AI AGENT 資料夾（外掛執行內嵌版本時傳入）")
    ap.add_argument("--building", default="")
    ap.add_argument("--level", default="")
    ap.add_argument("--grid", default="")
    ap.add_argument("--cellpos", default="", help="柱線格內相對位置 fx,fy（Revit 外掛自動帶入）")
    ap.add_argument("--point", default="", help="問題點位 X,Y（公尺，模型座標）")
    ap.add_argument("--colrefs", default="", help="周圍結構柱 編號:X,Y;…（公尺）")
    ap.add_argument("--labels", nargs="*", default=[], help="選取構件的圖說編號（C56、J7a）")
    ap.add_argument("--cats", nargs="*", default=[])
    ap.add_argument("--types", nargs="*", default=[])
    ap.add_argument("--keywords", nargs="*", default=[])
    ap.add_argument("--out", required=True)
    ap.add_argument("--top", type=int, default=4)
    a = ap.parse_args()
    if a.root:
        set_root(a.root)
    st = settings()
    index = st.get("index") or DEFAULT_INDEX
    pdf_dir = st.get("pdf_dir") or DEFAULT_PDF
    res = {"ok": False}
    try:
        if a.item:
            it = json.load(open(a.item, encoding="utf-8-sig"))
            a.building = a.building or it.get("Building", "")
            a.level = a.level or it.get("Level", "")
            a.grid = a.grid or it.get("Grid", "")
            a.cellpos = a.cellpos or it.get("CellPos", "")
            a.point = a.point or it.get("PointXY", "")
            a.colrefs = a.colrefs or it.get("ColumnRefs", "")
            a.labels = a.labels or [x for x in re.split(r"[、,;\s]+", it.get("Labels", "")) if x]
            els = it.get("Elements") or []
            a.cats = a.cats or [e.split(" ")[0] for e in els]
            a.types = a.types or els
            if not a.building:
                for b in CFG.get("buildings", []):
                    if b in (it.get("Location", "") + it.get("Question", "")):
                        a.building = b
        if not os.path.exists(index):
            raise SystemExit("找不到圖說查找索引：" + index)
        con = sqlite3.connect(index)
        meta = dict(con.execute("SELECT key,value FROM meta"))
        # 檢查 PDF 版本
        warn = []
        for f, size in con.execute("SELECT file,size FROM files"):
            p = os.path.join(pdf_dir, f)
            if not os.path.exists(p):
                warn.append(f"缺少 {f}")
            elif os.path.getsize(p) != size:
                warn.append(f"{f} 與索引的版本不同")
        refs = parse_refs(a.colrefs)
        labels = list(dict.fromkeys(label_key(x) for x in a.labels if label_key(x)))
        try:
            point = tuple(float(v) for v in a.point.split(",")) if a.point else None
        except ValueError:
            point = None
        cands, gparts, marks = find(con, a.building, a.level, a.grid, a.cats, a.types, a.keywords, a.top, refs, labels)
        os.makedirs(a.out, exist_ok=True)
        out = []
        for k, (score, pid, why) in enumerate(cands, 1):
            f, pg, no, title, gv, gh, kind = con.execute(
                "SELECT file,page,drawing_no,title,grid_v,grid_h,kind FROM pages WHERE id=?", (pid,)).fetchone()
            # 只有平面圖才依柱線裁切；立面/剖面/詳圖給整張
            gvl, ghl = json.loads(gv or "[]"), json.loads(gh or "[]")
            box = grid_box(gvl, ghl, gparts) if (gparts and kind == "平面圖") else None
            png = os.path.join(a.out, f"cand_{k}.png")
            path = os.path.join(pdf_dir, f)
            info = {}
            if os.path.exists(path):
                info = render(path, pg, box, png, gv=gvl if box else None, gh=ghl, gparts=gparts, cellpos=a.cellpos,
                              point=point, refs=refs, labels=labels, con=con, pid=pid, plan=(kind == "平面圖"))
            else:
                png = ""
            if info.get("method") == "柱編號":
                why.append("依柱編號定位（" + "、".join(info.get("fit_labels", [])[:5]) + f"，誤差 {info.get('fit_resid_m')}m）")
            elif info.get("snapped"):
                why.append("已依結構柱校正位置")
            if info.get("labels_found"):
                why.append("標出 " + "、".join(info["labels_found"]))
            located = bool(info.get("box")) if info else bool(box)
            out.append({"rank": k, "score": score, "why": why, "file": f, "page": pg, "drawing_no": no, "title": title,
                        "caption": f"{no} {title}".strip() + (f"（{a.grid}）" if located and a.grid else ""),
                        "png": png, "png_zoom": info.get("png_zoom", ""), "located": located, "marker": bool(info.get("marker")),
                        "method": info.get("method", ""), "snapped": bool(info.get("snapped")),
                        "column_check": info.get("column_check"), "fit_labels": info.get("fit_labels")})
        history = []
        try:
            history = find_history(a.building or guess_building(gparts), a.level, gparts)
        except Exception as e:
            warn.append(f"讀取釋疑歷史失敗：{e}")
        res = {"ok": True, "query": {"building": a.building, "level": a.level, "level_norm": norm_level(a.level),
                                     "grid": a.grid, "cellpos": a.cellpos, "point": a.point, "colrefs": len(refs),
                                     "labels": labels, "cats": a.cats, "marks": marks},
               "candidates": out, "history": history, "warnings": warn, "index_built": meta.get("built_at", ""),
               "index_builder": meta.get("builder", "")}
    except SystemExit as e:
        res["error"] = str(e)
    except Exception as e:
        import traceback
        res["error"] = f"{type(e).__name__}: {e}"
        res["trace"] = traceback.format_exc()[-1200:]
    os.makedirs(a.out, exist_ok=True)
    json.dump(res, open(os.path.join(a.out, "seek_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("RESULT " + json.dumps(res, ensure_ascii=False))


set_root(os.path.join(HERE, "..", ".."))

if __name__ == "__main__":
    main()
